from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook

from app.services.contract_review import service as contract_review_service_module
from app.services.contract_review.checklist_parser import parse_checklist_template
from app.services.contract_review.processor import ContractReviewProcessor
from app.services.loaders import registry
from app.services.loaders.base import DocumentLoadError, LoadResult, LoadedFragment
from app.services.loaders.word import WordDocumentLoader


def _build_checklist_workbook() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "2股权转让"
    rows = [
        ("适应范围", "1、适合结构较简单的股权转让合同。"),
        ("", "2、适合转移控制权的交易。"),
        ("风险名称", "转让方应当是目标公司的股东"),
        ("风险等级", "高"),
        ("风险描述", "1. 转让方一般应当是目标公司经工商登记的股东。"),
        ("", "2. 如存在代持，应当补充陈述与保证条款。"),
        ("相关知识点", "【1】隐名股东满足一定条件可以转让股权"),
        ("风险名称", "应当约定过渡期安排"),
        ("风险等级", "高"),
        ("风险描述", "1. 应限制过渡期内目标公司的异常经营行为。"),
        ("相关", "【2】过渡期内损害公司价值的责任承担"),
        ("知识点", "【3】过渡期违约责任的认定"),
        ("相关条款", "过渡期"),
    ]
    for row_index, (left, right) in enumerate(rows, start=1):
        sheet.cell(row=row_index, column=1, value=left)
        sheet.cell(row=row_index, column=2, value=right)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _inline_run_review_job(monkeypatch, *, llm_payload: dict | None = None) -> None:
    def run_inline(job_id: str) -> None:
        from app.core.config import get_settings
        from app.db.session import get_session_factory

        session = get_session_factory()()
        try:
            ContractReviewProcessor(session, get_settings()).run(job_id)
        finally:
            session.close()

    monkeypatch.setattr(contract_review_service_module, "submit_review_job", run_inline)

    if llm_payload is not None:
        monkeypatch.setattr(
            "app.services.contract_review.llm.ContractReviewLLMClient.is_configured",
            lambda self: True,
        )
        monkeypatch.setattr(
            "app.services.contract_review.llm.ContractReviewLLMClient.chat_json",
            lambda self, **_: llm_payload,
        )


def test_checklist_parser_supports_block_excel_structure(tmp_path) -> None:
    file_path = tmp_path / "template.xlsx"
    file_path.write_bytes(_build_checklist_workbook())

    parsed = parse_checklist_template(file_path)

    assert parsed.sheet_name == "2股权转让"
    assert parsed.applicability == [
        "1、适合结构较简单的股权转让合同。",
        "2、适合转移控制权的交易。",
    ]
    assert len(parsed.checklist) == 2
    assert parsed.checklist[0].title == "转让方应当是目标公司的股东"
    assert parsed.checklist[0].risk_level == "高"
    assert "代持" in parsed.checklist[0].description
    assert parsed.checklist[1].related_clauses == ["过渡期"]
    assert parsed.checklist[1].related_knowledge_points == [
        "【2】过渡期内损害公司价值的责任承担",
        "【3】过渡期违约责任的认定",
    ]


def test_contract_review_settings_roundtrip(client) -> None:
    response = client.get("/api/v1/contract-review/settings")
    assert response.status_code == 200
    assert response.json() == {"global_rule_prompt": ""}

    update_response = client.patch(
        "/api/v1/contract-review/settings",
        json={"global_rule_prompt": "所有审查结果都必须强调证据充分性。"},
    )
    assert update_response.status_code == 200
    assert update_response.json()["global_rule_prompt"].startswith("所有审查结果")

    second_read = client.get("/api/v1/contract-review/settings")
    assert second_read.status_code == 200
    assert second_read.json()["global_rule_prompt"].startswith("所有审查结果")


def test_template_upload_duplicate_review_and_export(client, monkeypatch) -> None:
    def fake_load_document(_file_path):
        return LoadResult(
            loader_name="WordDocumentLoader",
            fragments=[
                LoadedFragment(
                    fragment_index=0,
                    page_number=1,
                    content=(
                        "第一条 转让方资格\n"
                        "转让方系目标公司登记股东。\n\n"
                        "第二条 过渡期安排\n"
                        "交割前目标公司应保持正常经营，不得处分核心资产。\n"
                    ),
                    metadata={},
                )
            ],
        )

    monkeypatch.setattr(registry, "load_document", fake_load_document)
    _inline_run_review_job(
        monkeypatch,
        llm_payload={
            "items": [
                {
                    "checklist_key": "item-001",
                    "title": "转让方应当是目标公司的股东",
                    "severity": "high",
                    "status": "pass",
                    "issue": "合同已明确转让方为登记股东。",
                    "rewrite_suggestion": None,
                    "evidence_items": [
                        {
                            "clause_path": "C001",
                            "clause_title": "第一条 转让方资格",
                            "excerpt": "转让方系目标公司登记股东。",
                        }
                    ],
                },
                {
                    "checklist_key": "item-002",
                    "title": "应当约定过渡期安排",
                    "severity": "medium",
                    "status": "warn",
                    "issue": "已约定过渡期限制，但建议补充违约责任。",
                    "rewrite_suggestion": "补充过渡期内异常经营的赔偿责任和解除权。",
                    "evidence_items": [
                        {
                            "clause_path": "C002",
                            "clause_title": "第二条 过渡期安排",
                            "excerpt": "交割前目标公司应保持正常经营，不得处分核心资产。",
                        }
                    ],
                },
            ]
        },
    )

    template_bytes = _build_checklist_workbook()
    template_upload = client.post(
        "/api/v1/contract-review/templates/upload",
        data={"name": "股权转让清单"},
        files={
            "file": (
                "template.xlsx",
                template_bytes,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert template_upload.status_code == 200
    template_payload = template_upload.json()
    template_id = template_payload["items"][0]["id"]
    assert template_payload["items"][0]["config_json"]["template_mode"] == "xlsx_checklist"
    assert len(template_payload["items"][0]["config_json"]["checklist"]) == 2

    duplicate_upload = client.post(
        "/api/v1/contract-review/templates/upload",
        data={"name": "股权转让清单-重复"},
        files={
            "file": (
                "template.xlsx",
                template_bytes,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert duplicate_upload.status_code == 200
    assert len(duplicate_upload.json()["items"]) == 1

    settings_response = client.patch(
        "/api/v1/contract-review/settings",
        json={"global_rule_prompt": "所有 fail/warn 结果都应附修改建议。"},
    )
    assert settings_response.status_code == 200

    review_response = client.post(
        "/api/v1/contract-review/jobs",
        data={"template_id": template_id, "review_name": "股权转让协议审查"},
        files={
            "file": (
                "review.docx",
                b"review-bytes",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert review_response.status_code == 200
    payload = review_response.json()
    assert payload["job"]["status"] == "completed"
    assert payload["checklist"]["total"] == 2
    assert payload["checklist"]["passed"] == 1
    assert payload["checklist"]["warnings"] == 1
    assert payload["findings"][0]["checklist_item"]["risk_level"] == "高"
    assert payload["findings"][0]["evidence_items"][0]["clause_path"] == "C001"
    assert payload["findings"][1]["rewrite_suggestion"].startswith("补充过渡期")
    assert payload["export"]["available"] is True

    export_response = client.get(
        f"/api/v1/contract-review/jobs/{payload['job']['id']}/export.docx"
    )
    assert export_response.status_code == 200
    assert (
        export_response.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

    delete_response = client.delete(f"/api/v1/contract-review/templates/{template_id}")
    assert delete_response.status_code == 200
    assert delete_response.json()["items"] == []

    create_after_delete = client.post(
        "/api/v1/contract-review/jobs",
        data={"template_id": template_id},
        files={
            "file": (
                "review.docx",
                b"review-bytes",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert create_after_delete.status_code == 404


def test_review_export_conflict_before_completion(client, monkeypatch) -> None:
    monkeypatch.setattr(
        contract_review_service_module, "submit_review_job", lambda _job_id: None
    )

    template_upload = client.post(
        "/api/v1/contract-review/templates/upload",
        data={"name": "股权转让清单"},
        files={
            "file": (
                "template.xlsx",
                _build_checklist_workbook(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    template_id = template_upload.json()["items"][0]["id"]

    job_response = client.post(
        "/api/v1/contract-review/jobs",
        data={"template_id": template_id, "review_name": "待处理任务"},
        files={
            "file": (
                "review.docx",
                b"review-bytes",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert job_response.status_code == 200
    job_id = job_response.json()["job"]["id"]
    assert job_response.json()["job"]["status"] == "queued"

    export_response = client.get(f"/api/v1/contract-review/jobs/{job_id}/export.docx")
    assert export_response.status_code == 409


def test_contract_review_job_can_be_deleted_after_completion(
    client, monkeypatch
) -> None:
    def fake_load_document(_file_path):
        return LoadResult(
            loader_name="WordDocumentLoader",
            fragments=[
                LoadedFragment(
                    fragment_index=0,
                    page_number=1,
                    content="第一条 转让方资格\n转让方系目标公司登记股东。",
                    metadata={},
                )
            ],
        )

    monkeypatch.setattr(registry, "load_document", fake_load_document)
    _inline_run_review_job(
        monkeypatch,
        llm_payload={
            "items": [
                {
                    "checklist_key": "item-001",
                    "title": "转让方应当是目标公司的股东",
                    "severity": "high",
                    "status": "pass",
                    "issue": "条款已经明确。",
                    "rewrite_suggestion": None,
                    "evidence_items": [
                        {
                            "clause_path": "C001",
                            "clause_title": "第一条 转让方资格",
                            "excerpt": "转让方系目标公司登记股东。",
                        }
                    ],
                },
                {
                    "checklist_key": "item-002",
                    "title": "应当约定过渡期安排",
                    "severity": "medium",
                    "status": "missing",
                    "issue": "合同中未见过渡期条款。",
                    "rewrite_suggestion": "补充交割前过渡期行为限制。",
                    "evidence_items": [],
                },
            ]
        },
    )

    template_upload = client.post(
        "/api/v1/contract-review/templates/upload",
        data={"name": "删除任务模板"},
        files={
            "file": (
                "template.xlsx",
                _build_checklist_workbook(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    template_id = template_upload.json()["items"][0]["id"]

    job_response = client.post(
        "/api/v1/contract-review/jobs",
        data={"template_id": template_id, "review_name": "删除任务"},
        files={
            "file": (
                "review.docx",
                b"review-bytes",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert job_response.status_code == 200
    job_id = job_response.json()["job"]["id"]
    assert job_response.json()["job"]["status"] == "completed"

    delete_response = client.delete(f"/api/v1/contract-review/jobs/{job_id}")
    assert delete_response.status_code == 200
    assert delete_response.json()["affected_ids"] == [job_id]
    assert delete_response.json()["items"] == []

    detail_response = client.get(f"/api/v1/contract-review/jobs/{job_id}")
    assert detail_response.status_code == 404


def test_running_contract_review_job_cannot_be_deleted(client, monkeypatch) -> None:
    monkeypatch.setattr(
        contract_review_service_module, "submit_review_job", lambda _job_id: None
    )

    template_upload = client.post(
        "/api/v1/contract-review/templates/upload",
        data={"name": "运行中任务模板"},
        files={
            "file": (
                "template.xlsx",
                _build_checklist_workbook(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    template_id = template_upload.json()["items"][0]["id"]

    job_response = client.post(
        "/api/v1/contract-review/jobs",
        data={"template_id": template_id, "review_name": "运行中任务"},
        files={
            "file": (
                "review.docx",
                b"review-bytes",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert job_response.status_code == 200
    job_id = job_response.json()["job"]["id"]
    assert job_response.json()["job"]["status"] == "queued"

    delete_response = client.delete(f"/api/v1/contract-review/jobs/{job_id}")
    assert delete_response.status_code == 409


def test_contract_review_rejects_non_xlsx_template_upload(client) -> None:
    response = client.post(
        "/api/v1/contract-review/templates/upload",
        data={"name": "错误模板"},
        files={
            "file": (
                "template.docx",
                b"template-bytes",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 400
    assert ".xlsx" in response.json()["detail"]


def test_contract_review_rejects_legacy_doc_upload(client, monkeypatch) -> None:
    monkeypatch.setattr(
        contract_review_service_module, "submit_review_job", lambda _job_id: None
    )

    template_upload = client.post(
        "/api/v1/contract-review/templates/upload",
        data={"name": "股权转让清单"},
        files={
            "file": (
                "template.xlsx",
                _build_checklist_workbook(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    template_id = template_upload.json()["items"][0]["id"]

    response = client.post(
        "/api/v1/contract-review/jobs",
        data={"template_id": template_id, "review_name": "旧版文档审查"},
        files={"file": ("legacy.doc", b"legacy-doc-bytes", "application/msword")},
    )

    assert response.status_code == 400
    assert ".docx" in response.json()["detail"]


def test_xlsx_review_job_fails_without_answer_model(client, monkeypatch) -> None:
    def fake_load_document(_file_path):
        return LoadResult(
            loader_name="WordDocumentLoader",
            fragments=[
                LoadedFragment(
                    fragment_index=0,
                    page_number=1,
                    content="第一条 转让方资格\n转让方系目标公司登记股东。",
                    metadata={},
                )
            ],
        )

    monkeypatch.setattr(registry, "load_document", fake_load_document)
    _inline_run_review_job(monkeypatch)
    monkeypatch.setattr(
        "app.services.contract_review.llm.ContractReviewLLMClient.is_configured",
        lambda self: False,
    )

    template_upload = client.post(
        "/api/v1/contract-review/templates/upload",
        data={"name": "股权转让清单"},
        files={
            "file": (
                "template.xlsx",
                _build_checklist_workbook(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    template_id = template_upload.json()["items"][0]["id"]

    job_response = client.post(
        "/api/v1/contract-review/jobs",
        data={"template_id": template_id, "review_name": "模型未配置"},
        files={
            "file": (
                "review.docx",
                b"review-bytes",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert job_response.status_code == 200
    payload = job_response.json()
    assert payload["job"]["status"] == "failed"
    assert "模型未配置" in payload["job"]["failure_reason"]


def test_word_loader_reports_invalid_docx(tmp_path) -> None:
    file_path = tmp_path / "broken.docx"
    file_path.write_bytes(b"not-a-real-docx")

    loader = WordDocumentLoader()

    try:
        loader.load(file_path)
    except DocumentLoadError as exc:
        assert ".docx" in str(exc)
    else:
        raise AssertionError("Expected invalid .docx to raise DocumentLoadError")
