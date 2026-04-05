from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.models import ContractReviewClause, ContractReviewFinding, ContractReviewJob


def build_review_report(
    *,
    job: ContractReviewJob,
    clauses: list[ContractReviewClause],
    findings: list[ContractReviewFinding],
    output_path: Path,
) -> Path:
    document = Document()

    title = job.result_snapshot_json.get("overview", {}).get("title", "合同审查报告")
    heading = document.add_heading(title, level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = document.add_paragraph()
    meta.add_run(f"审查任务：{job.review_name}\n").bold = True
    meta.add_run(f"合同文件：{job.original_filename}\n")
    meta.add_run(f"任务状态：{job.status}\n")
    meta.add_run(
        f"模板：{job.result_snapshot_json.get('template_name', '未命名模板')}\n"
    )

    document.add_heading("一、审查摘要", level=1)
    overview = job.result_snapshot_json.get("overview", {})
    for highlight in overview.get("highlights", []):
        document.add_paragraph(str(highlight), style="List Bullet")

    document.add_heading("二、审查清单", level=1)
    checklist = job.summary_json
    table = document.add_table(rows=1, cols=4)
    header = table.rows[0].cells
    header[0].text = "总项数"
    header[1].text = "通过"
    header[2].text = "预警"
    header[3].text = "缺失/失败"
    row = table.add_row().cells
    row[0].text = str(checklist.get("total", 0))
    row[1].text = str(checklist.get("passed", 0))
    row[2].text = str(checklist.get("warnings", 0))
    row[3].text = str(checklist.get("missing", 0) + checklist.get("failed", 0))

    document.add_heading("三、风险与改写建议", level=1)
    if findings:
        for index, finding in enumerate(findings, start=1):
            paragraph = document.add_paragraph()
            paragraph.add_run(
                f"{index}. [{finding.severity.upper()} / {finding.status}] {finding.title}"
            ).bold = True
            document.add_paragraph(f"问题：{finding.issue}")
            if finding.evidence:
                document.add_paragraph(f"证据：{finding.evidence}")
            if finding.rewrite_suggestion:
                document.add_paragraph(f"建议改写：{finding.rewrite_suggestion}")
    else:
        document.add_paragraph("当前任务尚未生成风险项。")

    document.add_heading("四、条款预览", level=1)
    for clause in clauses[:12]:
        document.add_heading(
            f"{clause.clause_path} {clause.title}（第 {clause.page_start} 页）", level=2
        )
        document.add_paragraph(clause.content[:800])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path
