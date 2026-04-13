from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


class ChecklistTemplateParseError(ValueError):
    pass


@dataclass(slots=True)
class ParsedChecklistItem:
    key: str
    title: str
    risk_level: str
    description: str
    related_clauses: list[str]
    related_knowledge_points: list[str]
    source_block_text: str


@dataclass(slots=True)
class ParsedChecklistTemplate:
    sheet_name: str
    applicability: list[str]
    checklist: list[ParsedChecklistItem]


LABEL_ALIASES = {
    "适应范围": "applicability",
    "适用范围": "applicability",
    "风险名称": "title",
    "风险等级": "risk_level",
    "风险描述": "description",
    "相关条款": "related_clauses",
    "相关知识点": "related_knowledge_points",
    "相关": "related_knowledge_points",
    "知识点": "related_knowledge_points",
}


def parse_checklist_template(file_path: Path) -> ParsedChecklistTemplate:
    workbook = load_workbook(file_path, data_only=True, read_only=True)
    try:
        if not workbook.sheetnames:
            raise ChecklistTemplateParseError("审查清单中未找到工作表。")

        sheet = workbook[workbook.sheetnames[0]]
        sheet_name = str(sheet.title).strip() or "Sheet1"
        applicability: list[str] = []
        checklist: list[ParsedChecklistItem] = []
        current_item: dict[str, object] | None = None
        current_field: str | None = None

        def flush_item() -> None:
            nonlocal current_item
            if current_item is None:
                return
            title = str(current_item.get("title", "")).strip()
            if not title:
                raise ChecklistTemplateParseError("审查清单中存在缺少风险名称的区块。")

            checklist.append(
                ParsedChecklistItem(
                    key=str(current_item["key"]),
                    title=title,
                    risk_level=str(current_item.get("risk_level", "")).strip()
                    or "未标注",
                    description="\n".join(
                        line
                        for line in current_item.get("description_lines", [])
                        if isinstance(line, str) and line.strip()
                    ).strip(),
                    related_clauses=[
                        line
                        for line in current_item.get("related_clauses", [])
                        if isinstance(line, str) and line.strip()
                    ],
                    related_knowledge_points=[
                        line
                        for line in current_item.get("related_knowledge_points", [])
                        if isinstance(line, str) and line.strip()
                    ],
                    source_block_text="\n".join(
                        line
                        for line in current_item.get("source_lines", [])
                        if isinstance(line, str) and line.strip()
                    ).strip(),
                )
            )
            current_item = None

        def append_source_line(
            label: str | None,
            value: str,
            *,
            is_applicability: bool = False,
        ) -> None:
            if not value:
                return
            if is_applicability:
                prefix = "适用范围"
            else:
                prefix = label

            line = f"{prefix}: {value}" if prefix else value
            if is_applicability:
                applicability.append(value)
            elif current_item is not None:
                source_lines = current_item.setdefault("source_lines", [])
                assert isinstance(source_lines, list)
                source_lines.append(line)

        def append_item_value(field: str, value: str, label: str | None) -> None:
            if current_item is None or not value:
                return

            append_source_line(label, value)
            if field == "title":
                existing = str(current_item.get("title", "")).strip()
                current_item["title"] = (
                    f"{existing} {value}".strip() if existing else value
                )
                return
            if field == "risk_level":
                existing = str(current_item.get("risk_level", "")).strip()
                current_item["risk_level"] = (
                    f"{existing} {value}".strip() if existing else value
                )
                return
            if field == "description":
                description_lines = current_item.setdefault("description_lines", [])
                assert isinstance(description_lines, list)
                description_lines.append(value)
                return
            if field == "related_clauses":
                related_clauses = current_item.setdefault("related_clauses", [])
                assert isinstance(related_clauses, list)
                related_clauses.append(value)
                return
            if field == "related_knowledge_points":
                knowledge_points = current_item.setdefault(
                    "related_knowledge_points", []
                )
                assert isinstance(knowledge_points, list)
                knowledge_points.append(value)

        for row in sheet.iter_rows(values_only=True):
            left = _normalize_cell(row[0] if len(row) > 0 else None)
            right = _normalize_cell(row[1] if len(row) > 1 else None)

            if not left and not right:
                continue

            field = LABEL_ALIASES.get(left)
            if field == "title":
                flush_item()
                current_item = {
                    "key": f"item-{len(checklist) + 1:03d}",
                    "title": "",
                    "risk_level": "",
                    "description_lines": [],
                    "related_clauses": [],
                    "related_knowledge_points": [],
                    "source_lines": [],
                }
                current_field = field
                append_item_value(field, right, left)
                continue

            if current_item is None:
                if field == "applicability":
                    current_field = "applicability"
                    append_source_line(left, right, is_applicability=True)
                elif right and current_field == "applicability":
                    append_source_line(None, right, is_applicability=True)
                continue

            if field:
                current_field = field
                append_item_value(field, right, left)
                continue

            if right and current_field:
                append_item_value(current_field, right, None)

        flush_item()

        if not checklist:
            raise ChecklistTemplateParseError("未能从 Excel 清单中解析出任何风险项。")

        return ParsedChecklistTemplate(
            sheet_name=sheet_name,
            applicability=applicability,
            checklist=checklist,
        )
    finally:
        workbook.close()


def _normalize_cell(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r", "\n").strip()
    return text
